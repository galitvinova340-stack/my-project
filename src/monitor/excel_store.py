"""Запись карточек в общий Excel-файл на сетевом диске (раздел 3.4 PRD).

Если файл/папка недоступны (устройство выключено, диск не примонтирован),
запись не выполняется, но и не теряется — вызывающий код кладёт карточку
в локальную очередь (см. state.StateStore) и повторяет попытку в следующий
прогон, без потери данных и без дублирования.
"""
from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import EXCEL_HEADERS, AppointmentCard

logger = logging.getLogger(__name__)


class ExcelUnavailable(Exception):
    """Сетевой диск/файл недоступен в данный момент."""


def is_reachable(excel_dir: Path) -> bool:
    """Проверяет, смонтирован ли диск/сетевой ресурс, на котором должна лежать
    папка с таблицей. Саму папку (например, "назначения") мы вправе создать
    сами при первом запуске — недоступность диска определяется по его корню
    (букве диска или корню UNC-пути), а не по наличию конечной подпапки."""
    try:
        return Path(excel_dir.anchor).exists()
    except OSError:
        return False


def _open_or_create_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Назначения"
    ws.append(EXCEL_HEADERS)
    return wb


def read_existing_dedup_keys(path: Path) -> list[str]:
    """Читает уже внесённые записи (ФИО + должность) для сверки при дедупликации,
    в т.ч. строки, добавленные коллегами вручную."""
    from .models import normalize_key_part

    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    keys = []
    rows = ws.iter_rows(min_row=2, values_only=True)
    for row in rows:
        if not row or len(row) < 5:
            continue
        full_name, _photo, new_position = row[1], row[2], row[3]
        if full_name and new_position:
            keys.append(f"{normalize_key_part(str(full_name))}|{normalize_key_part(str(new_position))}")
    return keys


def append_card(excel_path: Path, card: AppointmentCard) -> None:
    """Добавляет одну строку в конец таблицы. Бросает ExcelUnavailable, если
    сетевая папка недоступна прямо сейчас."""
    if not is_reachable(excel_path.parent):
        raise ExcelUnavailable(f"Папка недоступна: {excel_path.parent}")

    try:
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb = _open_or_create_workbook(excel_path)
        ws = wb.active
        ws.append(card.as_excel_row())
        wb.save(excel_path)
    except OSError as exc:
        raise ExcelUnavailable(str(exc)) from exc

    logger.info("Карточка добавлена в Excel: %s — %s", card.full_name, card.new_position)
