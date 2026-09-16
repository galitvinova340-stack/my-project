from datetime import datetime

from monitor import excel_store
from monitor.models import EXCEL_HEADERS, AppointmentCard


def make_card(name="Иванов Иван Иванович") -> AppointmentCard:
    return AppointmentCard(
        fixed_at=datetime(2026, 9, 16, 8, 0),
        full_name=name,
        photo_url="",
        new_position="Аким района",
        organization="Акимат района",
        appointment_date="16.09.2026",
        previous_position="Замакима",
        career_path="",
        education="",
        contacts="",
        public_trace="",
        discussion_topics="",
        sources=["https://example.com"],
    )


def test_append_creates_file_with_headers(tmp_path):
    path = tmp_path / "sub" / "appointments.xlsx"
    excel_store.append_card(path, make_card())

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == EXCEL_HEADERS
    assert ws.max_row == 2


def test_append_is_additive_not_overwriting(tmp_path):
    path = tmp_path / "appointments.xlsx"
    excel_store.append_card(path, make_card("Иванов Иван"))
    excel_store.append_card(path, make_card("Петров Пётр"))

    from openpyxl import load_workbook

    ws = load_workbook(path).active
    assert ws.max_row == 3  # header + 2 rows


def test_read_existing_dedup_keys(tmp_path):
    path = tmp_path / "appointments.xlsx"
    card = make_card()
    excel_store.append_card(path, card)

    keys = excel_store.read_existing_dedup_keys(path)
    assert card.dedup_key in keys


def test_unavailable_drive_raises_and_does_not_write(tmp_path, monkeypatch):
    path = tmp_path / "appointments.xlsx"
    monkeypatch.setattr(excel_store, "is_reachable", lambda _dir: False)

    import pytest

    with pytest.raises(excel_store.ExcelUnavailable):
        excel_store.append_card(path, make_card())
    assert not path.exists()
